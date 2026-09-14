#!/usr/bin/env python3
"""Genera el informe del Examen Transversal (DOCX y PDF) a partir de docs/evidencias/.

Uso, desde la raiz del repositorio:
    pip install python-docx pillow
    python3 docs/informe/generar_informe.py

Requiere LibreOffice (soffice) y poppler-utils (pdftotext, pdfinfo, pdffonts).

Ninguna cifra del informe se escribe a mano: tests, comprobaciones e2e, requests de la coleccion Postman,
migraciones, operaciones de la API y codigos obtenidos se leen de los .txt de docs/evidencias/,
que guardan la salida literal de cada comando. El indice se calcula en pasadas: se convierte a
PDF, se busca en que pagina quedo cada titulo y se regenera con esos numeros hasta que no
cambian. Si falta una captura, su lugar queda marcado como PENDIENTE.
"""
import json
import re
import shutil
import subprocess
import sys
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

RAIZ = Path(__file__).resolve().parents[2]
EVIDENCIAS = RAIZ / "docs" / "evidencias"
INFORME = RAIZ / "docs" / "informe"
BUILD = INFORME / "build"
POSTMAN = RAIZ / "docs" / "api" / "postman" / "coleccion-laminas.postman_collection.json"
LOGO = INFORME / "logo-ipss.png"
NOMBRE = "EXT_2_LUCENA_GOMEZ_VALERIA_OTTON"

FUENTE = "Arial"
TAM_CUERPO = 12
TAM_TABLA = 10
TAM_PIE = 10
ANCHO_UTIL = 16.59  # cm: hoja carta (21,59 cm) menos dos margenes de 2,5 cm
ALTO_MAX = 5.0  # cm; las filas de 3 columnas usan 4,0. Acota el peor caso de cada captura
PAGINAS_MIN, PAGINAS_MAX = 5, 15  # con Anexo A (aclaracion recibida por el grupo el 13-09-2026; el enunciado dice 5 a 10)

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]

# Figuras del cuerpo (las marcadas con ★ en la guia), en el orden en que aparecen
FIGURAS = [
    ("01", "Proyecto en Spring Initializr"),
    ("02", "Dependencias del pom.xml"),
    ("03", "Configuración en application.properties"),
    ("05", "Contenedores de MySQL y de la API en ejecución"),
    ("10", "Diagrama entidad-relación del esquema creado por Flyway"),
    ("11", "Swagger UI con las operaciones agrupadas"),
    ("08", "Migraciones registradas en flyway_schema_history"),
    ("38", "Actualización con el encabezado X-Usuario"),
    ("41", "Historial de un álbum eliminado"),
    ("42", "Revisiones de Envers en revinfo y lamina_aud"),
    ("13", "Validación de datos (400)"),
    ("23", "Carga en lote (201)"),
    ("24", "Lote con conflictos (409)"),
    ("26", "Registrar obtenidas (200)"),
    ("30", "Láminas faltantes (200)"),
    ("31", "Láminas repetidas (200)"),
    ("32", "Resumen del álbum (200)"),
    ("33", "Subir foto (200)"),
    ("35", "Archivo que no es imagen (415)"),
    ("47", "Totales de ./mvnw verify, del script e2e y de Newman (Postman)"),
    # Anexo A.1: peticiones de la coleccion Postman sin figura en el cuerpo (pie = nombre de la request en la coleccion)
    ("12", None), ("14", None), ("15", None), ("16", None), ("17", None), ("18", None), ("19", None),
    ("20", None), ("21", None), ("22", None), ("25", None), ("27", None), ("28", None), ("29", None),
    ("34", None), ("36", None), ("37", None), ("39", None), ("40", None),
    # Anexo A.2: proyecto, base de datos y salidas completas de las suites
    ("04", "Estructura de paquetes del proyecto"),
    ("07", "Scripts de migración en db/migration"),
    ("43", "Validador personalizado NumerosUnicosValidator"),
    ("09", "Tablas creadas por Flyway en MySQL"),
    ("46", "Final de la salida de Newman (colección Postman)"),
    ("06", "Arranque de la API: Flyway aplica V1 a V5"),
    ("44", "Final del script pruebas/e2e.sh"),
    ("45", "Salida de ./mvnw verify"),
]
NUMERO_FIGURA = {ev: i + 1 for i, (ev, _) in enumerate(FIGURAS)}
# Capturas de terminal que ocupan una fila completa; el resto va siempre en la grilla, asi el
# conteo de paginas del borrador (con marcadores PENDIENTE) es el mismo que con las capturas reales
ANCHAS = {"05", "06", "08", "44", "45", "47"}

# Nombre de archivo de cada captura segun docs/informe/GUIA_INFORME_Y_CAPTURAS.md §3
ARCHIVOS = {
    "01": "01-initializr.png", "02": "02-pom-dependencias.png", "03": "03-application-properties.png",
    "05": "05-docker-compose-ps.png", "08": "08-flyway-schema-history.png", "10": "10-diagrama-er.png",
    "11": "11-swagger-ui.png", "12": "12-album-crear-201.png", "13": "13-album-validacion-400.png",
    "23": "23-lote-201.png", "24": "24-lote-conflictos-409.png", "26": "26-registrar-200.png",
    "30": "30-faltantes.png", "31": "31-repetidas.png", "32": "32-resumen.png", "33": "33-foto-subir-200.png",
    "35": "35-foto-falsa-415.png", "38": "38-auditoria-header.png", "41": "41-historial-album.png",
    "42": "42-revinfo-lamina-aud.png", "47": "47-resumen-suites.png",
    "04": "04-estructura-proyecto.png", "06": "06-arranque-api.png", "07": "07-carpeta-migraciones.png",
    "09": "09-tablas-mysql.png", "12": "12-album-crear-201.png", "14": "14-album-duplicado-409.png",
    "15": "15-album-listar-paginado.png", "16": "16-album-obtener-404.png", "17": "17-album-actualizar-200.png",
    "18": "18-lamina-crear-201.png", "19": "19-lamina-fuera-rango-409.png", "20": "20-lamina-listar-filtro.png",
    "21": "21-lamina-actualizar-200.png", "22": "22-lamina-eliminar-204.png", "25": "25-lote-repetidos-400.png",
    "27": "27-registrar-no-catalogado-409.png", "28": "28-cantidad-patch-200.png",
    "29": "29-cantidad-negativa-409.png", "34": "34-foto-ver-200.png", "36": "36-foto-grande-413.png",
    "37": "37-foto-eliminar-204.png", "39": "39-historial-lamina.png", "40": "40-album-eliminar-204.png",
    "43": "43-validador-numeros-unicos.png", "44": "44-script-e2e.png", "45": "45-mvnw-verify.png",
    "46": "46-newman.png",
}

# Descripcion de cada operacion; el conjunto se contrasta con /v3/api-docs (evidencia 11)
DESCRIPCIONES = {
    ("POST", "/api/albumes"): "Crear álbum",
    ("GET", "/api/albumes"): "Listar álbumes paginados",
    ("GET", "/api/albumes/{id}"): "Obtener un álbum",
    ("PUT", "/api/albumes/{id}"): "Actualizar un álbum",
    ("DELETE", "/api/albumes/{id}"): "Eliminar álbum, láminas y fotos",
    ("GET", "/api/albumes/{albumId}/resumen"): "Resumen y porcentaje de avance",
    ("GET", "/api/albumes/{id}/historial"): "Historial del álbum (Envers)",
    ("POST", "/api/albumes/{albumId}/laminas"): "Agregar una lámina",
    ("GET", "/api/albumes/{albumId}/laminas"): "Listar láminas (filtro ?tipo=)",
    ("POST", "/api/albumes/{albumId}/laminas/lote"): "Carga en lote, todo o nada",
    ("POST", "/api/albumes/{albumId}/laminas/registrar"): "Registrar obtenidas por número",
    ("GET", "/api/albumes/{albumId}/laminas/faltantes"): "Láminas faltantes",
    ("GET", "/api/albumes/{albumId}/laminas/repetidas"): "Repetidas y cantidad de repetidas",
    ("GET", "/api/laminas/{id}"): "Obtener una lámina",
    ("PUT", "/api/laminas/{id}"): "Actualizar una lámina",
    ("DELETE", "/api/laminas/{id}"): "Eliminar una lámina",
    ("PATCH", "/api/laminas/{id}/cantidad"): "Sumar o restar copias",
    ("POST", "/api/laminas/{id}/foto"): "Subir foto (multipart)",
    ("GET", "/api/laminas/{id}/foto"): "Ver la foto",
    ("DELETE", "/api/laminas/{id}/foto"): "Quitar la foto",
    ("GET", "/api/laminas/{id}/historial"): "Historial de la lámina (Envers)",
}


# ----------------------------------------------------------------------------------------------
# Lectura de evidencias
# ----------------------------------------------------------------------------------------------
def leer(nombre):
    ruta = EVIDENCIAS / nombre
    if not ruta.exists():
        sys.exit(f"Falta la evidencia {ruta}")
    return ruta.read_text(encoding="utf-8")


def buscar(patron, texto, origen):
    m = re.search(patron, texto, re.M)
    if not m:
        sys.exit(f"No se encontró /{patron}/ en {origen}")
    return m


def plantilla(metodo, ruta, operaciones):
    """Devuelve la ruta de OpenAPI que corresponde a una ruta concreta (con ids o variables)."""
    segmentos = ruta.strip("/").split("/")
    for m, p in operaciones:
        partes = p.strip("/").split("/")
        if m == metodo and len(partes) == len(segmentos) and all(
                a == b or a.startswith("{") for a, b in zip(partes, segmentos)):
            return p
    sys.exit(f"La ruta {metodo} {ruta} no coincide con ninguna operación de la API")


def sin_prefijo(ruta):
    return ruta[len("/api"):] if ruta.startswith("/api/") else ruta


def cargar_datos():
    d = {}

    t = leer("45-mvnw-verify.txt")
    total = buscar(r"Tests run: (\d+), Failures: (\d+), Errors: (\d+), Skipped: (\d+)$", t, "45")
    d["tests"] = int(total[1])
    d["fallos_tests"] = int(total[2]) + int(total[3])
    if "BUILD SUCCESS" not in t:
        sys.exit("45-mvnw-verify.txt no registra BUILD SUCCESS")
    por_clase = {m[2]: int(m[1]) for m in re.finditer(r"Tests run: (\d+),.* -- in [\w.]+\.(\w+)$", t, re.M)}
    d["t_unit"] = por_clase["ColeccionServiceImplTest"]
    d["t_web"] = por_clase["AlbumControllerTest"]
    d["t_int"] = por_clase["ColeccionLaminasApplicationTests"]

    t = leer("44-script-e2e.txt")
    m = buscar(r"RESULTADO: (\d+) OK, (\d+) FAIL", t, "44")
    d["e2e_ok"], d["e2e_fail"] = int(m[1]), int(m[2])
    m = buscar(r"^# Evidencia generada el (\d{4})-(\d{2})-(\d{2})", t, "44")
    d["fecha"] = f"{int(m[3])} de {MESES[int(m[2]) - 1]} de {m[1]}"
    d["_e2e"] = t

    t = leer("46-newman.txt")
    m = buscar(r"│\s+requests │\s+(\d+) │\s+(\d+) │", t, "46")
    d["col_req"], d["col_req_ok"] = int(m[1]), int(m[1]) - int(m[2])
    m = buscar(r"│\s+assertions │\s+(\d+) │\s+(\d+) │", t, "46")
    d["col_tests"], d["col_tests_ok"] = int(m[1]), int(m[1]) - int(m[2])
    # Node puede intercalar advertencias (stderr, con su propio "[DEP...]") entre la URL y el codigo:
    # se avanza, sin salir del bloque de la request, hasta el primer "[NNN "
    estados = {m[1]: m[2] for m in re.finditer(r"^↳ (\d{2}) [^\n]*\n\s+[A-Z]+ \S+[^↳]*?\[(\d{3}) ", t, re.M)}

    t = leer("08-flyway-schema-history.txt")
    filas = re.findall(r"^\|\s+\d+ \| (\d+)\s+\| (.+?)\s+\| V\S+\s+\| .+? \|\s+\d+ \|\s+(\d) \|$", t, re.M)
    d["migraciones"] = len(filas)
    if any(f[2] != "1" for f in filas):
        sys.exit("08-flyway-schema-history.txt tiene migraciones con success distinto de 1")

    t = leer("06-arranque-api.txt")
    d["arranque_s"] = buscar(r"Started ColeccionLaminasApplication in ([\d.]+) seconds", t, "06")[1].replace(".", ",")
    d["migraciones_log"] = int(buscar(r"Successfully applied (\d+) migrations", t, "06")[1])

    t = leer("05-docker-compose-ps.txt")
    if not re.search(r"^coleccion-mysql .*\(healthy\)", t, re.M) or not re.search(r"^coleccion-api .* Up ", t, re.M):
        sys.exit("05-docker-compose-ps.txt no muestra mysql (healthy) y api Up")

    t = leer("42-revinfo-lamina-aud.txt")
    filas = re.findall(r"^\|\s+(\d+) \| \S+ \S+ \| (\w+)\s+\|\s+\d+ \|\s+(\d) \|", t, re.M)
    d["rev_del"], d["rev_user"] = filas[0][0], filas[0][1]
    d["rev_n"] = sum(1 for f in filas if f[0] == d["rev_del"] and f[2] == "2")

    t = leer("11-openapi-operaciones.txt")
    d["_ops"] = re.findall(r"^(GET|POST|PUT|PATCH|DELETE) (/\S+)$", t, re.M)
    d["ops"] = len(d["_ops"])
    d["rutas"] = int(buscar(r"^rutas: (\d+)$", t, "11")[1])
    d["grupos"] = len(re.findall(r"'[^']+'", buscar(r"^tags: (.+)$", t, "11")[1]))
    if set(d["_ops"]) != set(DESCRIPCIONES):
        sys.exit(f"Las operaciones de la API no coinciden con DESCRIPCIONES: {set(d['_ops']) ^ set(DESCRIPCIONES)}")

    # Requests de la coleccion Postman (nombre, metodo, ruta y codigo esperado) + codigo obtenido con Newman
    d["_peticiones"] = []
    for carpeta in json.loads(POSTMAN.read_text(encoding="utf-8"))["item"]:
        for item in carpeta["item"]:
            nombre = buscar(r"^(\d{2}) (.+) \((\d{3})\)$", item["name"], item["name"])
            req = item["request"]
            ruta = "/" + "/".join(req["url"]["path"])
            query = "&".join(f"{q['key']}={q['value']}" for q in req["url"].get("query", []))
            obtenido = estados.get(nombre[1])
            if obtenido is None:
                sys.exit(f"La request {nombre[1]} no aparece en 46-newman.txt")
            d["_peticiones"].append((nombre[1], req["method"], plantilla(req["method"], ruta, d["_ops"]), query,
                                     nombre[2], nombre[3], obtenido))
    return d


def cobertura(d):
    """Operaciones sin caso exitoso (2xx) en la coleccion Postman, indicando si el script e2e las ejecuta con 2xx."""
    exitosas_coleccion = {(m, r) for _, m, r, _, _, _, obtenido in d["_peticiones"] if obtenido.startswith("2")}
    resultado = []
    for metodo, ruta in d["_ops"]:
        if (metodo, ruta) in exitosas_coleccion:
            continue
        patron = re.sub(r"\\\{\w+\\\}", r"[^/ ?]+", re.escape(ruta))
        e2e = re.search(rf"^\[OK\] .*{metodo} {patron}(\?\S*)? -> 2\d\d", d["_e2e"], re.M) is not None
        resultado.append((metodo, ruta, e2e))
    return resultado


# ----------------------------------------------------------------------------------------------
# Formato
# ----------------------------------------------------------------------------------------------
def fuente_estilo(estilo, tam, negrita=False):
    estilo.font.name = FUENTE
    estilo.font.size = Pt(tam)
    estilo.font.bold = negrita
    estilo.font.italic = False
    rpr = estilo.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    for atributo in list(rfonts.attrib):  # quita las fuentes del tema (asciiTheme, etc.)
        del rfonts.attrib[atributo]
    for clave in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(clave), FUENTE)
    color = rpr.find(qn("w:color"))
    if color is not None:
        rpr.remove(color)
    estilo.font.color.rgb = RGBColor(0, 0, 0)


def formato_parrafo(pf, alineacion, antes=0, despues=4):
    pf.alignment = alineacion
    pf.line_spacing = 1.15
    pf.space_before = Pt(antes)
    pf.space_after = Pt(despues)


def configurar(doc):
    normal = doc.styles["Normal"]
    fuente_estilo(normal, TAM_CUERPO)
    formato_parrafo(normal.paragraph_format, WD_ALIGN_PARAGRAPH.JUSTIFY)
    idioma = OxmlElement("w:lang")
    idioma.set(qn("w:val"), "es-CL")
    normal.element.get_or_add_rPr().append(idioma)

    for nivel, tam in ((1, 14), (2, 12)):
        estilo = doc.styles[f"Heading {nivel}"]
        fuente_estilo(estilo, tam, negrita=True)
        formato_parrafo(estilo.paragraph_format, WD_ALIGN_PARAGRAPH.LEFT, antes=8 if nivel == 1 else 4)
        estilo.paragraph_format.keep_with_next = True

    try:
        pie = doc.styles["Caption"]
    except KeyError:
        pie = doc.styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
    fuente_estilo(pie, TAM_PIE)
    formato_parrafo(pie.paragraph_format, WD_ALIGN_PARAGRAPH.CENTER, antes=2, despues=6)

    seccion = doc.sections[0]
    seccion.page_width, seccion.page_height = Cm(21.59), Cm(27.94)
    for margen in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(seccion, margen, Cm(2.5))
    seccion.different_first_page_header_footer = True  # la portada no lleva numero
    pie_pagina = seccion.footer.paragraphs[0]
    pie_pagina.alignment = WD_ALIGN_PARAGRAPH.CENTER
    campo_pagina(pie_pagina)


def campo_pagina(parrafo):
    for tipo in ("begin", "instr", "separate", "texto", "end"):
        run = parrafo.add_run()
        if tipo == "texto":
            run.text = "1"
            continue
        if tipo == "instr":
            elemento = OxmlElement("w:instrText")
            elemento.set(qn("xml:space"), "preserve")
            elemento.text = "PAGE"
        else:
            elemento = OxmlElement("w:fldChar")
            elemento.set(qn("w:fldCharType"), tipo)
        run._r.append(elemento)


def sombrear(celda, color):
    sombra = OxmlElement("w:shd")
    sombra.set(qn("w:val"), "clear")
    sombra.set(qn("w:color"), "auto")
    sombra.set(qn("w:fill"), color)
    celda._tc.get_or_add_tcPr().append(sombra)


def fijar_anchos(tabla, anchos):
    """Anchos de columna en la grilla (w:tblGrid) y en cada celda: LibreOffice usa la grilla."""
    tabla.autofit = False
    diseno = OxmlElement("w:tblLayout")
    diseno.set(qn("w:type"), "fixed")
    tabla._tbl.tblPr.append(diseno)
    for columna, ancho in zip(tabla._tbl.tblGrid.findall(qn("w:gridCol")), anchos):
        columna.set(qn("w:w"), str(Cm(ancho).twips))
    for fila in tabla.rows:
        for celda, ancho in zip(fila.cells, anchos):
            celda.width = Cm(ancho)


def propiedad_fila(fila, nombre):
    elemento = OxmlElement(nombre)
    elemento.set(qn("w:val"), "true")
    fila._tr.get_or_add_trPr().append(elemento)


def normalizar(texto):
    return " ".join(unicodedata.normalize("NFKC", texto).split())


# ----------------------------------------------------------------------------------------------
# Documento
# ----------------------------------------------------------------------------------------------
class Informe:
    def __init__(self, datos, paginas):
        self.doc = Document()
        self.d = datos
        self.paginas = paginas or {}
        self.encabezados = []
        self.faltantes = []
        self.figuras_puestas = 0
        self.tablas = 0
        configurar(self.doc)

    # --- texto ---
    def sustituir(self, texto):
        def reemplazo(m):
            clave = m[1]
            if clave.startswith("fig:"):
                return f"Figura {NUMERO_FIGURA[clave[4:]]}"
            return str(self.d[clave])
        return re.sub(r"<<([\w:]+)>>", reemplazo, texto)

    def texto(self, parrafo, texto, tam=None, negrita=False):
        for parte in re.split(r"(\*[^*]+\*)", self.sustituir(texto)):
            if not parte:
                continue
            cursiva = len(parte) > 2 and parte[0] == "*" and parte[-1] == "*"
            run = parrafo.add_run(parte[1:-1] if cursiva else parte)
            run.italic = cursiva or None
            run.bold = negrita or None
            if tam:
                run.font.size = Pt(tam)
        return parrafo

    def p(self, texto):
        return self.texto(self.doc.add_paragraph(), texto)

    def titulo(self, nivel, texto, salto=False):
        parrafo = self.doc.add_heading(texto, level=nivel)
        parrafo.paragraph_format.page_break_before = salto
        self.encabezados.append((nivel, texto))

    # --- tablas ---
    def tabla(self, titulo, encabezado, filas, anchos, centradas=()):
        self.tablas += 1
        leyenda = self.texto(self.doc.add_paragraph(style="Caption"), f"Tabla {self.tablas}: {titulo}")
        leyenda.alignment = WD_ALIGN_PARAGRAPH.LEFT
        leyenda.paragraph_format.keep_with_next = True
        tabla = self.doc.add_table(rows=1, cols=len(encabezado))
        tabla.style = "Table Grid"
        tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
        tabla.autofit = False
        propiedad_fila(tabla.rows[0], "w:tblHeader")
        for i, valor in enumerate(encabezado):
            self.celda(tabla.rows[0].cells[i], valor, negrita=True, centrada=True)
            sombrear(tabla.rows[0].cells[i], "D9D9D9")
        for fila in filas:
            nueva = tabla.add_row()
            propiedad_fila(nueva, "w:cantSplit")
            for i, valor in enumerate(fila):
                self.celda(nueva.cells[i], valor, centrada=i in centradas)
        fijar_anchos(tabla, anchos)
        separador = self.doc.add_paragraph()
        separador.paragraph_format.space_after = Pt(0)

    def celda(self, celda, valor, negrita=False, centrada=False):
        parrafo = celda.paragraphs[0]
        formato_parrafo(parrafo.paragraph_format,
                        WD_ALIGN_PARAGRAPH.CENTER if centrada else WD_ALIGN_PARAGRAPH.LEFT, despues=0)
        self.texto(parrafo, str(valor), tam=TAM_TABLA, negrita=negrita)

    # --- figuras ---
    def ruta(self, ev):
        exacta = EVIDENCIAS / ARCHIVOS[ev]
        if exacta.exists():
            return exacta
        otras = sorted(EVIDENCIAS.glob(f"{ev}-*.png"))
        return otras[0] if otras else None

    def ancha(self, ev):
        return ev in ANCHAS

    def imagen(self, parrafo, ev, ancho_max, alto_max):
        ruta = self.ruta(ev)
        parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if ruta is None:
            self.faltantes.append(ARCHIVOS[ev])
            # reserva el alto maximo: el conteo de paginas del borrador es el peor caso
            parrafo.paragraph_format.space_before = Cm(alto_max / 2 - 0.3)
            parrafo.paragraph_format.space_after = Cm(alto_max / 2 - 0.3)
            run = parrafo.add_run(f"[PENDIENTE: {ARCHIVOS[ev]}]")
            run.bold = True
            run.font.size = Pt(TAM_PIE)
            run.font.color.rgb = RGBColor(0xC0, 0, 0)
            return
        w, h = Image.open(ruta).size
        parrafo.add_run().add_picture(str(ruta), width=Cm(min(ancho_max, alto_max * w / h)))

    def leyenda(self, parrafo, ev):
        esperado = FIGURAS[self.figuras_puestas]
        if esperado[0] != ev:
            sys.exit(f"Figura {ev} fuera de orden: se esperaba la {esperado[0]}")
        self.figuras_puestas += 1
        parrafo.style = self.doc.styles["Caption"]
        pie = esperado[1] or next(f"{caso} ({codigo})" for n, _, _, _, caso, codigo, _ in self.d["_peticiones"] if n == ev)
        self.texto(parrafo, f"Figura {NUMERO_FIGURA[ev]}: {pie} (ev. {ev}).")

    def figuras(self, *evs, columnas=2, alto_max=ALTO_MAX):
        """Ubica las figuras en filas de `columnas`; las muy anchas (terminal) ocupan una fila completa."""
        fila = []
        for ev in evs:
            if self.ancha(ev):
                if fila:
                    self.fila(fila, columnas, alto_max)
                    fila = []
                self.fila([ev], 1, alto_max)
                continue
            fila.append(ev)
            if len(fila) == columnas:
                self.fila(fila, columnas, alto_max)
                fila = []
        if fila:
            self.fila(fila, columnas, alto_max)

    def fila(self, evs, columnas, alto_max):
        ancho = ANCHO_UTIL / columnas
        tabla = self.doc.add_table(rows=1, cols=len(evs))
        tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
        fijar_anchos(tabla, [ancho] * len(evs))
        propiedad_fila(tabla.rows[0], "w:cantSplit")
        for celda, ev in zip(tabla.rows[0].cells, evs):
            imagen = celda.paragraphs[0]
            formato_parrafo(imagen.paragraph_format, WD_ALIGN_PARAGRAPH.CENTER, despues=0)
            self.imagen(imagen, ev, ancho - 0.3, alto_max)
            self.leyenda(celda.add_paragraph(), ev)
        # Parrafo de 1 pt entre filas: Word y LibreOffice fusionan tablas contiguas y la grilla combinada
        # de filas con distinto numero de celdas deforma los anchos
        separador = self.doc.add_paragraph()
        formato_parrafo(separador.paragraph_format, WD_ALIGN_PARAGRAPH.LEFT, despues=0)
        separador.paragraph_format.line_spacing = Pt(1)

    # --- secciones ---
    def portada(self):
        doc = self.doc
        logo = doc.add_paragraph()
        logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo.paragraph_format.space_before = Cm(1)
        logo.add_run().add_picture(str(LOGO), width=Cm(8))
        lineas = [
            ("Sistema de Gestión de Colección de Láminas", 20, True, Cm(4.5)),
            ("API REST con Spring Boot y MySQL", 14, False, Pt(6)),
            ("Examen Transversal", 14, False, Pt(18)),
            ("Asignatura: Desarrollo de Software Web II", TAM_CUERPO, False, Cm(5)),
            ("Profesor: Boris Marcelo Belmar Muñoz", TAM_CUERPO, False, Pt(6)),
            ("Integrantes: Valeria Gómez y Otton Lucena", TAM_CUERPO, False, Pt(6)),
            ("Grupo 2", TAM_CUERPO, False, Pt(6)),
            ("Septiembre de 2026", TAM_CUERPO, False, Pt(6)),
        ]
        for texto, tam, negrita, antes in lineas:
            parrafo = doc.add_paragraph()
            parrafo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            parrafo.paragraph_format.space_before = antes
            self.texto(parrafo, texto, tam=tam, negrita=negrita)

    def indice(self):
        titulo = self.doc.add_paragraph()
        titulo.paragraph_format.page_break_before = True
        titulo.alignment = WD_ALIGN_PARAGRAPH.LEFT
        self.texto(titulo, "Índice", tam=14, negrita=True)
        return self.doc.add_paragraph()  # marcador: las entradas se insertan antes al final

    def completar_indice(self, marcador):
        for nivel, texto in self.encabezados:
            entrada = marcador.insert_paragraph_before()
            formato_parrafo(entrada.paragraph_format, WD_ALIGN_PARAGRAPH.LEFT, despues=3)
            entrada.paragraph_format.left_indent = Cm(0 if nivel == 1 else 0.8)
            entrada.paragraph_format.tab_stops.add_tab_stop(Cm(ANCHO_UTIL), WD_TAB_ALIGNMENT.RIGHT,
                                                            WD_TAB_LEADER.DOTS)
            run = entrada.add_run(f"{texto}\t{self.paginas.get(texto, 0)}")
            run.bold = nivel == 1 or None
        marcador._element.getparent().remove(marcador._element)

    def introduccion(self):
        self.titulo(1, "1. Introducción", salto=True)
        self.p("Quienes coleccionan láminas necesitan saber cuáles les faltan y cuáles tienen repetidas para "
               "decidir qué comprar e intercambiar, algo difícil de controlar a mano cuando el álbum tiene "
               "cientos de láminas. Nuestro objetivo fue construir una API REST en Spring Boot y MySQL que "
               "administre álbumes y láminas, permita cargarlas en lote y registrar las obtenidas, y entregue las "
               "faltantes, las repetidas con su cantidad de repetidas y un resumen del avance, validando los "
               "datos y registrando cada cambio. El alcance es el que fija el enunciado, un backend REST sin "
               "frontend obligatorio; la autenticación de usuarios quedó fuera.")
        self.titulo(2, "1.1 Justificación de la tecnología")
        self.p("Usamos Spring Boot 4.1.1 y Java 21, versión de soporte extendido, porque el enunciado pide un "
               "proyecto Maven creado con Spring Initializr (Spring, s. f.-c) y porque Spring Web MVC, Spring "
               "Data JPA y Bean Validation cubren los requisitos sin infraestructura propia (Spring, s. f.-a). "
               "MySQL 8.4, también de soporte extendido, es obligatoria según el enunciado y el profesor, y "
               "corresponde a un problema relacional: cada lámina pertenece a un álbum y su número no se repite "
               "en él, reglas que MySQL garantiza con claves foráneas, UNIQUE y CHECK (Oracle Corporation, s. f.).")
        self.p("Para la auditoría seguimos la indicación del profesor de usar migraciones: Flyway 12.4 aplica "
               "scripts SQL versionados y registra cada cambio del esquema en flyway_schema_history (Redgate "
               "Software, s. f.), y JPA Auditing con Hibernate Envers 7.4 registran quién y cuándo cambió cada "
               "registro, incluso eliminado (Hibernate, s. f.). Docker Compose levanta todo con un comando "
               "(Docker Inc., s. f.); springdoc-openapi 3.1, la línea compatible con Spring Boot 4, documenta la "
               "API (springdoc-openapi, s. f.); Postman guarda las pruebas en una colección JSON versionada, "
               "que también corre en la terminal con Newman (Postman, s. f.), y Testcontainers prueba contra "
               "MySQL 8.4 real en lugar de una base "
               "en memoria (Testcontainers, s. f.).")

    def configuracion(self):
        self.titulo(1, "2. Desarrollo")
        self.titulo(2, "2.1 Configuración del proyecto y conexión a la base de datos")
        self.p("Generamos el proyecto en Spring Initializr con Maven, Java 21, Spring Boot 4.1.1, grupo cl.ipss, "
               "artefacto coleccion-laminas y las dependencias Spring Web, Spring Data JPA, MySQL Driver, "
               "Validation, Flyway Migration y Lombok (<<fig:01>>). Al pom.xml agregamos a mano solo "
               "hibernate-envers, con la versión que administra Spring Boot, y springdoc-openapi 3.1.0; "
               "flyway-mysql habilita a Flyway para MySQL y las dependencias de prueba tienen alcance test "
               "(<<fig:02>>).")
        self.p("El archivo application.properties toma la conexión de las variables DB_URL, DB_USER y "
               "DB_PASSWORD, con valores por defecto de desarrollo; así el mismo archivo sirve en el equipo local "
               "y en Docker, y las credenciales quedan en un .env que no se versiona (<<fig:03>>). Con "
               "ddl-auto=validate, Hibernate no toca el esquema y solo verifica al arrancar que coincide con las "
               "entidades; con open-in-view=false ninguna consulta ocurre fuera de los servicios.")
        self.p("El archivo compose.yaml define el servicio mysql (imagen mysql:8.4, volumen mysql-data y "
               "healthcheck) y el servicio api, construido con un Dockerfile de dos etapas (JDK 21 para compilar "
               "y JRE 21 con un usuario sin privilegios para ejecutar), que arranca solo cuando MySQL está sano. "
               "En la verificación final del <<fecha>>, hecha desde cero con volúmenes nuevos y docker compose up "
               "-d --build, MySQL quedó healthy y la API activa (<<fig:05>>); Flyway aplicó <<migraciones_log>> migraciones y "
               "la aplicación arrancó en <<arranque_s>> segundos (<<fig:06>>).")
        self.figuras("01", "02", "03", "05", columnas=3, alto_max=4.0)

    def arquitectura(self):
        self.titulo(2, "2.2 Arquitectura y modelo de datos")
        self.p("La aplicación se organiza en las capas controller, service y repository (<<fig:04>>). Los cinco "
               "controladores delegan en un servicio, sin lógica de negocio ni acceso a repositorios, y los tres "
               "que reciben un cuerpo JSON (álbumes, láminas y colección) lo validan con @Valid. Cada servicio "
               "tiene interfaz e implementación, "
               "concentra las reglas de negocio y define las transacciones (@Transactional en las escrituras y "
               "readOnly en las lecturas); los repositorios extienden JpaRepository (Spring, s. f.-b). Entre "
               "capas viajan DTOs definidos como record y convertidos por mappers manuales, así ninguna entidad "
               "JPA se expone en la API.")
        self.p("El enunciado pide una API REST y no exige una interfaz gráfica, por lo que no desarrollamos un "
               "frontend. La interfaz de usuario del sistema es Swagger UI (/swagger-ui.html), que muestra y "
               "permite ejecutar las <<ops>> operaciones con sus parámetros, ejemplos y códigos de respuesta, "
               "junto con la colección de Postman de <<col_req>> peticiones. Para quien la usa, todos los errores "
               "llegan en español y con un formato único.")
        self.p("Album guarda nombre, que es único, imagen, fecha de lanzamiento, tipo de láminas, total, "
               "editorial y descripción; Lamina guarda número, nombre, tipo, cantidad de copias y el archivo de "
               "su foto opcional (<<fig:10>>). La relación muchos a uno, con carga diferida, es la clave foránea "
               "album_id con ON DELETE CASCADE, y UNIQUE (album_id, numero) impide repetir un número en el álbum. "
               "Ambas entidades heredan de EntidadAuditable (creado_en, creado_por, modificado_en y "
               "modificado_por) y el historial se guarda en revinfo, album_aud y lamina_aud. El estado de una "
               "lámina sale de un solo campo, cantidad: 0 es faltante, 1 o más es obtenida y más de 1 es "
               "repetida, con cantidad − 1 repetidas. Un contador, en vez de tablas separadas, evita estados "
               "inconsistentes, y CHECK (cantidad >= 0) lo protege en la base.")
        self.figuras("10", "11")

    def api(self):
        self.titulo(2, "2.3 API REST y funcionalidades especiales")
        self.p("La API expone <<ops>> operaciones sobre <<rutas>> rutas, todas bajo el prefijo /api (Tabla 1). "
               "Crear un álbum o una lámina responde 201 con el encabezado Location y la carga en lote, 201 con "
               "las láminas creadas; las lecturas y actualizaciones responden 200 y las eliminaciones 204. El "
               "listado de álbumes es paginado, con un máximo de 100 elementos por "
               "página, y Swagger UI agrupa las operaciones en <<grupos>> secciones (<<fig:11>>).")
        orden = list(DESCRIPCIONES)
        filas = sorted(((m, sin_prefijo(r), DESCRIPCIONES[(m, r)]) for m, r in self.d["_ops"]),
                       key=lambda f: orden.index((f[0], "/api" + f[1])))
        self.tabla("Operaciones de la API, sin el prefijo /api (fuente: /v3/api-docs, evidencia 11)",
                   ["Método", "Ruta", "Función"], filas, [1.8, 7.4, 7.4])
        self.p('La carga en lote es todo o nada: @NumerosUnicos, nuestro validador personalizado, rechaza con '
               '400 los números repetidos en la lista; luego el servicio consulta una vez los números ya '
               'catalogados y, si alguna posición excede el total o ya existe, responde 409 con los errores por '
               'posición (laminas[0].numero, laminas[1].numero) sin guardar nada; si no hay conflictos, guarda '
               'todo con saveAll dentro de una transacción. El registro de obtenidas recibe números como '
               '{"numeros": [1, 5, 5, 12]}, busca las láminas con una consulta IN y suma una copia por '
               'aparición, o rechaza todo con 409 si hay números no catalogados; el PATCH de cantidad aplica un '
               'delta sin bajar de cero.')
        self.p("Faltantes y repetidas se resuelven en la base de datos, no en memoria: dos métodos derivados de "
               "Spring Data en LaminaRepository generan consultas con WHERE cantidad = 0 y WHERE cantidad > 1, "
               "y cada repetida informa su cantidad de repetidas. Las faltantes incluyen además los números aún "
               "no catalogados, calculados a partir de una consulta que trae solo la columna numero, y el "
               "resumen sale de una única consulta JPQL con count y sum(case ...). Filtrar en Java obligaría a "
               "cargar todas las láminas del álbum en cada consulta, mientras que MySQL usa el índice de UNIQUE "
               "(album_id, numero) y devuelve solo las filas pedidas.")
        self.p("La foto se sube como multipart/form-data, en el campo archivo, y se guarda en el volumen uploads "
               "con un nombre UUID; la base guarda solo ese nombre. Descartamos Base64 y BLOB porque Base64 "
               "aumenta el tamaño en un tercio y engrosa filas y respaldos, y el profesor lo indicó como mala "
               "práctica. El tipo se valida por los primeros bytes (JPEG FF D8 FF, PNG 89 50 4E 47, WEBP "
               "RIFF...WEBP) y no por la extensión, el máximo es 5 MB, la ruta se normaliza contra path "
               "traversal y los archivos se borran solo tras el commit, o tras el rollback si la subida falla.")

    def validaciones(self):
        self.titulo(2, "2.4 Validaciones y auditoría")
        self.p("Validamos en tres niveles: Bean Validation en los DTOs (@NotBlank, @NotNull, @NotEmpty, @Size, "
               "@Min, @Max, @PastOrPresent y @Pattern), activada con @Valid (Bean Validation, s. f.), junto con "
               "el validador personalizado @NumerosUnicos (@Constraint y ConstraintValidator), que responden "
               "400; reglas de negocio en los servicios, que responden 409 (nombre de álbum único, número no "
               "mayor que el total y único en el álbum, total no menor que la lámina más alta y ajustes que "
               "dejarían la cantidad bajo cero), y NOT NULL, UNIQUE, "
               "FOREIGN KEY y CHECK en la base. Un único GlobalExceptionHandler responde los errores 400, 404, "
               "409, 413, 415 y 500 con el formato ProblemDetail del RFC 9457 (Nottingham et al., 2023), con "
               "los errores por campo y sin exponer stack traces.")
        self.p("La auditoría tiene tres capas. La principal, pedida por el profesor, son las migraciones: el "
               "esquema se crea solo con Flyway (V1 álbum, V2 lámina, V3 columnas de auditoría, V4 tablas de "
               "Envers y V5 datos de ejemplo), queda versionado en Git y cada cambio se registra en "
               "flyway_schema_history con versión, checksum, fecha, duración y resultado; las <<migraciones>> "
               "migraciones figuran con success = 1 (<<fig:08>>). Como Flyway verifica el checksum de lo ya "
               "aplicado, alterar una migración impide arrancar la aplicación, así que el historial del esquema "
               "no se puede reescribir.")
        self.p("JPA Auditing guarda en cada registro quién y cuándo lo creó y modificó, con el usuario del "
               "encabezado X-Usuario o «sistema» si falta o no es válido; la <<fig:38>> muestra un cambio de «ana» "
               "sobre un álbum creado por «equipo». Hibernate Envers guarda cada versión en album_aud y "
               "lamina_aud, y en revinfo la revisión, la fecha y el usuario, en la misma transacción del cambio, "
               "de modo que el historial se actualiza en tiempo real. Con store_data_at_delete la eliminación "
               "conserva los datos, y las láminas de un álbum borrado se eliminan como entidades para que cada "
               "una quede auditada. Los endpoints /historial exponen cada revisión con fecha, usuario y operación "
               "(<<fig:41>>), y en MySQL la revisión <<rev_del>> de «<<rev_user>>» registra con revtype 2 "
               "(eliminación) las <<rev_n>> láminas del álbum eliminado (<<fig:42>>). Cada escritura deja además "
               "una línea INFO en el log.")
        self.figuras("08", "38", "41", "42", columnas=3, alto_max=4.0)

    def pruebas(self):
        d = self.d
        self.titulo(1, "3. Pruebas")
        self.p("Probamos en tres niveles el <<fecha>>, sobre un entorno recién creado: ./mvnw verify ejecutó "
               "<<tests>> pruebas con <<fallos_tests>> fallos (<<t_unit>> unitarias con Mockito, <<t_web>> de la "
               "capa web con @WebMvcTest y <<t_int>> de integración con Testcontainers, que aplica las migraciones "
               "en MySQL 8.4 y valida el esquema); pruebas/e2e.sh hizo <<e2e_ok>> comprobaciones con "
               "<<e2e_fail>> fallos, incluidas consultas a flyway_schema_history y lamina_aud, y Newman, el ejecutor de línea de comandos de Postman, aprobó "
               "<<col_req_ok>> de <<col_req>> peticiones y <<col_tests_ok>> de <<col_tests>> aserciones "
               "(<<fig:47>>; las salidas completas están en el Anexo A).")
        sin_exito = cobertura(d)
        en_e2e = [f"{m} {r}" for m, r, e2e in sin_exito if e2e]
        sin_prueba = [f"{m} {r}" for m, r, e2e in sin_exito if not e2e]
        nota = ""
        if en_e2e:
            nota += (f"El caso exitoso de {' y '.join(en_e2e)} no tiene una petición propia en la colección; "
                     "lo ejecuta el script end-to-end con respuesta 2xx (<<fig:44>>). ")
        if sin_prueba:
            nota += (f"El caso exitoso de {' y '.join(sin_prueba)} no quedó cubierto por ninguna de las dos "
                     "suites, que solo prueban su caso de error.")
        self.p("La Tabla 2 muestra, para cada petición de la colección de Postman, el código esperado, el obtenido con Newman y la "
               "figura con su captura: las más representativas están a continuación y el resto en el Anexo A. "
               + nota)
        filas = []
        for numero, metodo, ruta, query, caso, esperado, obtenido in d["_peticiones"]:
            referencia = f"Fig. {NUMERO_FIGURA[numero]}" if numero in NUMERO_FIGURA else f"Ev. {numero}"
            filas.append((numero, f"{metodo} {sin_prefijo(ruta)}" + (f"?{query}" if query else ""), caso,
                          esperado, obtenido, referencia))
        self.tabla("Peticiones de la colección de Postman: código esperado y obtenido con Newman (fuente: evidencia 46)",
                   ["N°", "Petición", "Caso", "Esp.", "Obt.", "Evid."], filas,
                   [0.8, 7.3, 4.6, 1.1, 1.1, 1.7], centradas=(0, 3, 4, 5))
        self.figuras("13", "23", "24", "26", "30", "31", "32", "33", "35", columnas=3, alto_max=4.0)
        self.figuras("47", alto_max=3.4)

    def conclusion(self):
        self.titulo(1, "4. Conclusión")
        self.p("Cumplimos los requerimientos del enunciado: la API se conecta a MySQL 8.4 en Docker, ofrece el "
               "CRUD con foto opcional, la carga en lote y los listados de faltantes y repetidas calculados en la "
               "base, separa la lógica de negocio del acceso a datos, valida en tres niveles y audita con "
               "migraciones, JPA Auditing y Envers; tres suites de pruebas lo verifican.")
        self.p("Resolvimos tres problemas: los enum que Hibernate 7 mapeaba al tipo ENUM de MySQL, corregidos "
               "con @JdbcTypeCode(SqlTypes.VARCHAR); las láminas borradas por la cascada de la base, que Envers "
               "no auditaba y ahora se eliminan como entidades, y las fechas de auditoría, truncadas a "
               "microsegundos por DATETIME(6). Quedan como mejoras la autenticación, pues X-Usuario identifica "
               "pero no autentica, y un frontend web.")

    def bibliografia(self):
        self.titulo(1, "5. Bibliografía")
        referencias = [
            "Bean Validation. (s. f.). *Jakarta Bean Validation 3.1*. https://beanvalidation.org/3.1/",
            "Docker Inc. (s. f.). *Docker Compose*. https://docs.docker.com/compose/",
            "Hibernate. (s. f.). *Hibernate ORM documentation* (Versión 7.4). "
            "https://hibernate.org/orm/documentation/",
            "Nottingham, M., Wilde, E. y Dalal, S. (2023). *Problem details for HTTP APIs* (RFC 9457). "
            "Internet Engineering Task Force. https://www.rfc-editor.org/rfc/rfc9457",
            "Oracle Corporation. (s. f.). *MySQL 8.4 reference manual*. https://dev.mysql.com/doc/refman/8.4/en/",
            "Postman. (s. f.). *Postman Learning Center*. https://learning.postman.com/",
            "Redgate Software. (s. f.). *Flyway documentation*. https://documentation.red-gate.com/flyway",
            "Spring. (s. f.-a). *Spring Boot reference documentation* (Versión 4.1). "
            "https://docs.spring.io/spring-boot/",
            "Spring. (s. f.-b). *Spring Data JPA reference documentation*. "
            "https://docs.spring.io/spring-data/jpa/reference/",
            "Spring. (s. f.-c). *Spring Initializr*. https://start.spring.io/",
            "springdoc-openapi. (s. f.). *springdoc-openapi* (Versión 3.1). https://springdoc.org/",
            "Testcontainers. (s. f.). *Testcontainers for Java*. https://java.testcontainers.org/",
        ]
        for referencia in referencias:
            parrafo = self.p(referencia)
            parrafo.alignment = WD_ALIGN_PARAGRAPH.LEFT  # APA 7: referencias alineadas a la izquierda
            parrafo.paragraph_format.space_after = Pt(0)
            parrafo.paragraph_format.left_indent = Cm(1.27)
            parrafo.paragraph_format.first_line_indent = Cm(-1.27)

    def anexo(self):
        self.titulo(1, "Anexo A. Capturas de las pruebas", salto=True)
        self.p("Este anexo completa la evidencia del capítulo 3: reúne la captura de cada petición de la Tabla 2 "
               "que no aparece en el cuerpo, junto con el proyecto, la base de datos y las salidas completas de "
               "las suites automáticas. Las imágenes están a resolución completa y se pueden ampliar en el PDF.")
        self.titulo(2, "A.1 Peticiones de la colección de Postman")
        self.figuras(*[ev for ev, pie in FIGURAS if pie is None], columnas=3)
        self.titulo(2, "A.2 Proyecto, base de datos y suites automáticas")
        self.figuras("04", "07", "43", "09", "46", columnas=3)
        self.figuras("06", "44", "45")

    def construir(self):
        self.portada()
        marcador = self.indice()
        self.introduccion()
        self.configuracion()
        self.arquitectura()
        self.api()
        self.validaciones()
        self.pruebas()
        self.conclusion()
        self.bibliografia()
        self.anexo()
        if self.figuras_puestas != len(FIGURAS):
            sys.exit(f"Se ubicaron {self.figuras_puestas} de {len(FIGURAS)} figuras")
        self.completar_indice(marcador)
        return self


# ----------------------------------------------------------------------------------------------
# PDF e indice
# ----------------------------------------------------------------------------------------------
def a_pdf(docx):
    perfil = (BUILD / "perfil-libreoffice").resolve()
    subprocess.run(["soffice", f"-env:UserInstallation=file://{perfil}", "--headless", "--convert-to", "pdf",
                    "--outdir", str(BUILD), str(docx)], check=True, capture_output=True, timeout=300)
    return BUILD / f"{docx.stem}.pdf"


def paginas_pdf(pdf):
    salida = subprocess.run(["pdfinfo", str(pdf)], check=True, capture_output=True, text=True).stdout
    return int(re.search(r"^Pages:\s+(\d+)", salida, re.M)[1])


def ubicar(pdf, encabezados):
    total = paginas_pdf(pdf)
    textos = {i: normalizar(subprocess.run(["pdftotext", "-f", str(i), "-l", str(i), str(pdf), "-"],
                                           check=True, capture_output=True, text=True).stdout)
              for i in range(1, total + 1)}
    paginas, desde = {}, 3  # 1 = portada, 2 = indice
    for _, titulo in encabezados:
        clave = normalizar(titulo)
        encontrada = next((i for i in range(desde, total + 1) if clave in textos[i]), None)
        if encontrada is None:
            sys.exit(f"No se encontró el título «{titulo}» en el PDF")
        paginas[titulo] = desde = encontrada
    return paginas


def main():
    datos = cargar_datos()
    BUILD.mkdir(exist_ok=True)
    paginas = None
    for _ in range(4):
        informe = Informe(datos, paginas).construir()
        docx = BUILD / f"{NOMBRE}.docx"
        informe.doc.save(docx)
        pdf = a_pdf(docx)
        nuevas = ubicar(pdf, informe.encabezados)
        if nuevas == paginas:
            break
        paginas = nuevas
    else:
        sys.exit("El índice no se estabilizó")

    shutil.copy(docx, INFORME / docx.name)
    shutil.copy(pdf, INFORME / pdf.name)
    total = paginas_pdf(pdf)
    fuentes = subprocess.run(["pdffonts", str(pdf)], capture_output=True, text=True).stdout
    print(f"PDF: {INFORME / pdf.name}")
    print(f"Páginas: {total} (rango exigido {PAGINAS_MIN}-{PAGINAS_MAX})")
    for titulo, pagina in paginas.items():
        print(f"  {pagina:>2}  {titulo}")
    print("Fuentes del PDF:\n" + fuentes)
    # Sin Arial instalada, LibreOffice la sustituye (Liberation Sans) al exportar: el enunciado exige Arial 12
    if "Arial" not in fuentes or "LiberationSans" in fuentes:
        sys.exit("El PDF no incrusta Arial: LibreOffice la reemplazo por otra fuente. Instala Arial y vuelve a generar:\n"
                 "  echo 'ttf-mscorefonts-installer msttcorefonts/accepted-mscorefonts-eula select true'"
                 " | sudo debconf-set-selections\n"
                 "  sudo apt-get install -y ttf-mscorefonts-installer && fc-cache -f\n"
                 "  fc-match Arial    # debe responder Arial.ttf, no LiberationSans")
    if informe.faltantes:
        print(f"CAPTURAS PENDIENTES ({len(informe.faltantes)}): " + ", ".join(informe.faltantes))
    if not PAGINAS_MIN <= total <= PAGINAS_MAX:
        sys.exit(f"El PDF tiene {total} páginas, fuera del rango {PAGINAS_MIN}-{PAGINAS_MAX}")


if __name__ == "__main__":
    main()
